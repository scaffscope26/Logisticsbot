import csv, io, json, sqlite3, datetime, hashlib, math, zipfile
from pathlib import Path
from openpyxl import load_workbook
DB=Path(__file__).parent/'logistics.db'
SCHEMAS={
 'inventory':['facility','part_number','description','quantity_available'],
 'projects':['project_id','customer','facility','part_number','quantity_on_rent','expected_return_date'],
 'reservations':['reservation_id','facility','part_number','quantity_required','required_date'],
 'returns':['project_id','facility','part_number','quantity_returning','return_date']}
ALIASES={'branch':'facility','yard':'facility','location':'facility','part':'part_number','item':'part_number','item_number':'part_number','qty':'quantity_available','available':'quantity_available','job':'project_id','job_number':'project_id','order_id':'reservation_id','order_number':'reservation_id','ship_date':'required_date','delivery_date':'required_date','off_rent_date':'expected_return_date','pickup_date':'return_date'}
def norm(s):return str(s or '').strip().lower().replace(' ','_').replace('-','_').replace('/','_')
def db():
 c=sqlite3.connect(DB);c.row_factory=sqlite3.Row
 c.executescript('CREATE TABLE IF NOT EXISTS records (kind TEXT NOT NULL, record_key TEXT NOT NULL, payload TEXT NOT NULL, batch_id TEXT NOT NULL, PRIMARY KEY(kind,record_key));CREATE TABLE IF NOT EXISTS imports (batch_id TEXT PRIMARY KEY,kind TEXT,filename TEXT,created_at TEXT,rows INTEGER);CREATE TABLE IF NOT EXISTS decisions (opportunity_id TEXT PRIMARY KEY,status TEXT,owner TEXT,note TEXT,updated_at TEXT);')
 return c
def parse(data,filename):
 if filename.lower().endswith('.csv'):
  reader=csv.DictReader(io.StringIO(data.decode('utf-8-sig')))
  headers=reader.fieldnames
  check_headers(headers)
  result=[]
  for row in reader:
   if len(result)>=20000:raise ValueError('Reports are limited to 20,000 rows')
   if None in row:raise ValueError('A row has more values than the header')
   result.append(row)
  return result
 if filename.lower().endswith('.xlsx'):
  with zipfile.ZipFile(io.BytesIO(data)) as archive:
   if sum(info.file_size for info in archive.infolist())>40_000_000:raise ValueError('Expanded workbook exceeds 40 MB')
  wb=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
  try:
   ws=wb.active
   if ws.max_column and ws.max_column>200:raise ValueError('Reports are limited to 200 columns')
   if ws.max_row and ws.max_row>20001:raise ValueError('Reports are limited to 20,000 rows')
   rows=ws.iter_rows(values_only=True);headers=next(rows,None)
   check_headers(headers)
   result=[]
   for row in rows:
    if len(result)>=20000:raise ValueError('Reports are limited to 20,000 rows')
    if any(v is not None for v in row):result.append(dict(zip(headers,row)))
   return result
  finally:wb.close()
 raise ValueError('Only .csv and .xlsx are supported')
def check_headers(headers):
 if not headers or len(headers)>200 or any(not norm(h) for h in headers):raise ValueError('Use a single header row with named columns (200 maximum)')
 if len({norm(h) for h in headers})!=len(headers):raise ValueError('Column headings must be unique')
def canonical(row,mapping):
 r={};normalized={norm(k):v for k,v in row.items()}
 for target in mapping:
  src=norm(mapping[target]);r[target]=normalized.get(src,'')
 for key,value in normalized.items():
  t=ALIASES.get(key,key)
  if t not in r:r[t]=value
 return r
def validate_import(kind,data,filename,mapping):
 if kind not in SCHEMAS:raise ValueError('Unknown import type')
 rows=parse(data,filename)
 if not rows:raise ValueError('File has no data rows')
 fields=SCHEMAS[kind];missing=[f for f in fields if not mapping.get(f)]
 if missing:raise ValueError('Map required fields: '+', '.join(missing))
 batch=hashlib.sha256(data+kind.encode()).hexdigest()[:16]
 clean=[];errors=[]
 for i,row in enumerate(rows,2):
  r=canonical(row,mapping)
  if any(r.get(f) is None or str(r.get(f)).strip()=='' for f in fields):errors.append(f'Row {i}: missing required values');continue
  try:
   for k in ['quantity_available','quantity_on_rent','quantity_required','quantity_returning']:
    if k in r:
     r[k]=float(r[k])
     if not math.isfinite(r[k]) or r[k]<0:raise ValueError('Invalid quantity')
   for k in ['expected_return_date','required_date','return_date']:
    if k in r:
     v=r[k];r[k]=v.date().isoformat() if isinstance(v,datetime.datetime) else v.isoformat() if isinstance(v,datetime.date) else str(v).strip()[:10]
     datetime.date.fromisoformat(r[k])
  except Exception:errors.append(f'Row {i}: invalid quantity or date (use YYYY-MM-DD)');continue
  for field in fields:
   if not field.startswith('quantity_') and not field.endswith('_date'):
    r[field]=str(r[field]).strip()
  r['part_number']=r['part_number'].upper()
  clean.append({field:r[field] for field in fields})
 keys=[record_identity(kind,r) for r in clean]
 if len(set(keys))!=len(keys):errors.append('Duplicate material lines: consolidate quantities for each facility, part, and project/reservation before import')
 if errors:return {'ok':False,'errors':errors[:30],'accepted':len(clean),'rejected':len(errors)}
 return {'ok':True,'rows':clean,'accepted':len(clean),'rejected':0}

def ingest(kind,data,filename,mapping):
 result=validate_import(kind,data,filename,mapping)
 if not result['ok']:return result
 clean=result['rows'];batch=hashlib.sha256(data+kind.encode()).hexdigest()[:16]
 c=db()
 with c:
  # Each import replaces the current snapshot for its type, not other types.
  c.execute('DELETE FROM records WHERE kind=?',(kind,))
  for idx,r in enumerate(clean):
   key=str(idx)+'-'+hashlib.sha1(json.dumps(r,sort_keys=True,default=str).encode()).hexdigest()[:12]
   c.execute('INSERT INTO records VALUES (?,?,?,?)',(kind,key,json.dumps(r,default=str),batch))
  c.execute('INSERT OR REPLACE INTO imports VALUES (?,?,?,?,?)',(batch,kind,filename,datetime.datetime.now(datetime.timezone.utc).isoformat(),len(clean)))
 return {'ok':True,'accepted':len(clean),'rejected':0,'batch':batch}
def records(kind):
 c=db();out=[json.loads(x['payload']) for x in c.execute('SELECT payload FROM records WHERE kind=?',(kind,))];c.close();return out
def record_identity(kind,row):
 fields=['facility','part_number']
 if kind in ['projects','returns']:fields.insert(0,'project_id')
 if kind=='reservations':fields.insert(0,'reservation_id')
 return hashlib.sha256(json.dumps([kind]+[str(row[f]).strip() for f in fields]).encode()).hexdigest()[:24]

def calculate(snapshot, today=None):
 today=today or datetime.date.today();end=today+datetime.timedelta(days=30)
 inv=snapshot.get('inventory',[]);returns=snapshot.get('returns',[]);projects=snapshot.get('projects',[]);demands=snapshot.get('reservations',[])
 supply=[]
 for i,r in enumerate(inv):supply.append({'id':record_identity('inventory',r),'part':str(r['part_number']).strip().upper(),'facility':r['facility'],'qty':float(r['quantity_available']),'date':today,'type':'yard','confidence':'confirmed','source':r})
 for i,r in enumerate(returns):
  date=datetime.date.fromisoformat(r['return_date']);
  if today<=date<=end:supply.append({'id':record_identity('returns',r),'part':str(r['part_number']).strip().upper(),'facility':r['facility'],'qty':float(r['quantity_returning']),'date':date+datetime.timedelta(days=2),'type':'return','confidence':'pending inspection','source':r})
 # Projects are potential leads only; never allocate them as guaranteed supply.
 forecast=[]
 for r in projects:
  date=datetime.date.fromisoformat(r['expected_return_date'])
  if today<=date<=end:forecast.append({'facility':r['facility'],'part':r['part_number'],'qty':r['quantity_on_rent'],'date':date.isoformat(),'project':r['project_id'],'confidence':'unconfirmed'})
 supply.sort(key=lambda x:(x['date'],x['facility'],x['id']))
 results=[];remaining={s['id']:s['qty'] for s in supply}
 for i,d in enumerate(sorted(demands,key=lambda r:(r['required_date'],record_identity('reservations',r)))):
  due=datetime.date.fromisoformat(d['required_date']);need=float(d['quantity_required']);
  if due<today or due>end:continue
  part=str(d['part_number']).strip().upper();candidates=[s for s in supply if s['part']==part and s['date']<=due and remaining[s['id']]>0]
  # Same facility first; cross-facility is a proposal requiring transport validation.
  candidates.sort(key=lambda s:(s['facility']!=d['facility'],s['date']))
  for s in candidates:
   if need<=0:break
   qty=min(need,remaining[s['id']]);remaining[s['id']]-=qty;need-=qty
   oid=hashlib.sha1(f"{record_identity('reservations',d)}|{s['id']}".encode()).hexdigest()[:14]
   results.append({'id':oid,'reservation':d['reservation_id'],'part':part,'quantity':qty,'from':s['facility'],'to':d['facility'],'needed':due.isoformat(),'available':s['date'].isoformat(),'source_type':s['type'],'confidence':s['confidence'],'cross_branch':s['facility']!=d['facility'],'status':'Proposed','warning':'Confirm transport and inspection before approval' if s['facility']!=d['facility'] or s['type']=='return' else 'Confirm reservation remains open'})
  if need>0:results.append({'id':'short-'+str(i),'reservation':d['reservation_id'],'part':part,'quantity':need,'from':'Unmatched','to':d['facility'],'needed':due.isoformat(),'available':'','source_type':'shortage','confidence':'unmatched','cross_branch':False,'status':'Shortage','warning':'Review forecast projects and purchasing options'})
 return {'opportunities':results,'forecast':forecast,'summary':{'inventory_rows':len(inv),'project_rows':len(projects),'reservation_rows':len(demands),'return_rows':len(returns),'matches':sum(r['source_type']!='shortage' for r in results),'shortages':sum(r['source_type']=='shortage' for r in results)}}
def matches():
 result=calculate({kind:records(kind) for kind in SCHEMAS})
 c=db();decisions={r['opportunity_id']:dict(r) for r in c.execute('SELECT * FROM decisions')};c.close()
 for row in result['opportunities']:
  saved=decisions.get(row['id'])
  if saved:row.update({key:saved[key] for key in ['status','owner','note']})
 return result

def decision(oid,status,owner,note):
 if status not in ['Proposed','Investigating','Approved (internal only)','Rejected','Completed']:raise ValueError('Invalid status')
 c=db()
 with c:c.execute('INSERT OR REPLACE INTO decisions VALUES (?,?,?,?,?)',(oid,status,owner,note,datetime.datetime.now(datetime.timezone.utc).isoformat()))
